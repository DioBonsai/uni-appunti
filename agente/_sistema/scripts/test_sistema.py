"""Collaudo automatico delle parti nuove del sistema: sottolineature, pagine dei riassunti, sito locale.

Uso:  python test_sistema.py      (esce con errore se qualcosa non va; funziona anche con pytest)
Non tocca i Word veri: lavora su file temporanei.
"""
import http.client
import json
import re
import sys
import tempfile
import threading
from http.server import ThreadingHTTPServer
from pathlib import Path

HERE = Path(__file__).resolve().parent
BASE = HERE.parent
PAG = BASE / "pagina_studio"
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(PAG))


def test_applica_sottolineature():
    from docx import Document

    import applica_sottolineature as A

    d = Document()
    d.add_paragraph("Capitolo uno", style="Heading 1")
    d.add_paragraph("Un insieme è una collezione ben determinata di oggetti distinti.")
    d.add_paragraph("Un insieme  vuoto   non ha elementi; un insieme finito ne ha un numero finito.")
    p = d.add_paragraph()
    p.add_run("Già ").font.highlight_color = A.WD_COLOR_INDEX.YELLOW
    p.add_run("evidenziato davvero")
    with tempfile.TemporaryDirectory() as t:
        f = Path(t) / "prova.docx"
        d.save(str(f))
        prima = [q.text for q in Document(str(f)).paragraphs]
        voci = [
            {"id": "a", "m": "x", "q": "collezione ben determinata", "p": "Un insieme è una ", "s": " di oggetti", "h": "Capitolo uno"},
            {"id": "b", "m": "x", "q": "insieme vuoto non ha", "p": "", "s": "", "h": ""},   # spazi diversi nel Word
            {"id": "c", "m": "x", "q": "testo che non esiste", "p": "", "s": "", "h": ""},
            {"id": "d", "m": "x", "q": "Già evidenziato", "p": "", "s": "", "h": ""},
            {"id": "e", "m": "x", "q": "un numero finito", "p": "ne ha ", "s": ".", "h": ""},
        ]
        ok, gia, no, _ = A.applica_a_word(f, voci, False)
        assert set(ok) == {"a", "b", "d", "e"}, ok                  # «Già evidenziato» era evidenziato a metà: si completa il resto
        assert gia == [] and no == ["c"], (gia, no)
        dopo = Document(str(f))
        assert [q.text for q in dopo.paragraphs] == prima, "il testo non deve cambiare"
        ev = "".join(r.text for q in dopo.paragraphs for r in q.runs if r.font.highlight_color)
        assert "collezione ben determinata" in ev and "insieme  vuoto   non ha" in ev and "un numero finito" in ev and "Già evidenziato" in ev, ev
        ok2, gia2, no2, _ = A.applica_a_word(f, voci[:2], False)   # idempotente: la seconda volta è già tutto evidenziato
        assert ok2 == [] and set(gia2) == {"a", "b"}, (ok2, gia2)


def test_pagine_riassunti():
    cartelle = sorted((PAG / "riassunti").glob("*/index.html"))
    assert len(cartelle) == 5, f"attese 5 pagine, trovate {len(cartelle)}"
    for f in cartelle:
        s = f.read_text(encoding="utf-8")
        nav = s[s.index('<nav id="indice"'):s.index("</nav>")]
        art = s[s.index("<article>"):]
        ids = set(re.findall(r'\bid="([^"]+)"', art))
        mancanti = [h for h in re.findall(r'href="#([^"]+)"', nav) if h not in ids]
        assert not mancanti, f"{f.parent.name}: voci d'indice senza sezione: {mancanti[:3]}"
        for src in re.findall(r'<img src="([^"]+)"', art):
            assert (f.parent / src).is_file(), f"{f.parent.name}: figura mancante {src}"


def test_quaderno_costruito():
    s = (PAG / "index.html").read_text(encoding="utf-8")
    for pezzo in ('id="tab-riassunti"', "riass-lettore", "api/sottolineature"):
        assert pezzo in s, f"manca {pezzo} in index.html (rilancia costruisci_pagina_studio.py)"
    json.loads((PAG / "data.json").read_text(encoding="utf-8"))


def test_server_sottolineature():
    import serve

    tmp = Path(tempfile.mkdtemp()) / "sott.json"
    serve.SOTT = tmp
    srv = ThreadingHTTPServer(("127.0.0.1", 0), serve.make_handler(0))
    port = srv.server_address[1]
    srv.RequestHandlerClass = serve.make_handler(port)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        def chiama(metodo, path, body=None, host=None):
            c = http.client.HTTPConnection("127.0.0.1", port)
            c.putrequest(metodo, path, skip_host=True)
            c.putheader("Host", host or f"127.0.0.1:{port}")
            data = json.dumps(body).encode() if body is not None else b""
            if body is not None:
                c.putheader("Content-Type", "application/json")
                c.putheader("Content-Length", str(len(data)))
            c.endheaders(data)
            r = c.getresponse()
            return r.status, r.read()

        voce = {"id": "u1", "m": "algebra-lineare", "q": "testo", "p": "", "s": "", "h": "", "t": "2026-10-20"}
        assert chiama("GET", "/api/sottolineature")[0] == 200
        st, b = chiama("POST", "/api/sottolineature", {"add": [voce], "del": []})
        assert st == 200 and json.loads(b)["items"][0]["id"] == "u1"
        voce2 = dict(voce, id="u2")
        st, b = chiama("POST", "/api/sottolineature", {"add": [voce2], "del": ["u1"]})
        assert [i["id"] for i in json.loads(b)["items"]] == ["u2"]
        assert chiama("POST", "/api/sottolineature", {"add": [{"id": 5}], "del": []})[0] == 200   # voci malformate ignorate
        assert chiama("POST", "/api/sottolineature", {"add": [], "del": []}, host="evil.example")[0] == 403   # protezione dal DNS rebinding
        assert chiama("GET", "/riassunti/algebra-lineare/index.html")[0] == 200
        assert chiama("GET", "/riassunti/../serve.py")[0] in (403, 404)
    finally:
        srv.shutdown()


if __name__ == "__main__":
    falliti = 0
    for nome, fn in list(globals().items()):
        if nome.startswith("test_") and callable(fn):
            try:
                fn()
                print(f"[  OK  ] {nome}")
            except Exception as e:  # noqa: BLE001
                falliti += 1
                print(f"[ERRORE] {nome}: {type(e).__name__}: {e}")
    sys.exit(1 if falliti else 0)
