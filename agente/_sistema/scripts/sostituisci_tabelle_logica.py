"""Sostituisce nel Word di Fondamenti le tavole di verità scritte a simboli (righe Consolas) con le immagini di figure_logica.py.
Una tantum, su richiesta di Emanuele (14 ott 2026). Backup automatico, controllo che il resto del testo e le evidenziazioni restino identici.
Uso: python sostituisci_tabelle_logica.py [--dry-run]
"""
import shutil
import sys
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

BASE = Path(__file__).resolve().parent.parent
UNI = BASE.parent.parent
DOC = UNI / "Riassunto Fondamenti Informatica.docx"

# (inizio del primo paragrafo della tavola, numero di paragrafi da togliere compresa la riga vuota finale, figura, larghezza cm, didascalia)
BLOCCHI = [
    ("p q | p∧q  p∨q  xor  equiv", 10, "figure/fond/logica_connettivi.png", 13.0, "Le tavole di verità di and, or, xor, equivalenza e not (1 = vero, 0 = falso)."),
    ("p q | p∨q  ¬p | (p∨q)∧¬p", 6, "figure/fond/logica_formula.png", 10.5, "La tavola di (p ∨ q) ∧ ¬p, calcolata a passi: una colonna per ogni sottoformula."),
    ("p q | p∨q  ¬(p∨q) | ¬p  ¬q  ¬p∧¬q", 6, "figure/fond/logica_demorgan.png", 12.5, "Le colonne di ¬(p ∨ q) e di ¬p ∧ ¬q coincidono: le due formule sono equivalenti."),
]


def firma(doc, escludi):
    testi, evid = [], 0
    for p in doc.paragraphs:
        if p._p in escludi or p._p.xpath(".//w:drawing"):
            continue
        testi.append(p.text)
        for r in p.runs:
            rpr = r._r.rPr
            if rpr is not None and (rpr.find(qn("w:highlight")) is not None or rpr.find(qn("w:shd")) is not None):
                evid += 1
    return testi, evid


def main():
    dry = "--dry-run" in sys.argv
    lock = DOC.with_name("~$" + DOC.name[2:])
    if lock.exists() and not dry:
        sys.exit(f"Il Word e' aperto ({lock.name}). Chiudilo e riprova.")
    doc = Document(str(DOC))
    P = doc.paragraphs
    togli, nuovi = [], []   # elementi da eliminare / paragrafi aggiunti
    piano = []
    for inizio, n, fig, larg, did in BLOCCHI:
        idx = [i for i, p in enumerate(P) if p.text == inizio]
        if len(idx) != 1:
            sys.exit(f"Tavola non trovata in modo univoco: {inizio!r} ({len(idx)})")
        i = idx[0]
        blocco = P[i:i + n]
        if not all(all(r.font.name == "Consolas" for r in p.runs) for p in blocco) or blocco[-1].text.strip():
            sys.exit(f"Blocco inatteso a partire da {inizio!r}: controllo manuale")
        if any(r.font.highlight_color is not None for p in blocco for r in p.runs):
            sys.exit(f"Il blocco {inizio!r} contiene evidenziazioni: non lo tocco")
        piano.append((i, n, fig, larg, did))
    if dry:
        for i, n, fig, *_ in piano:
            print(f"paragrafi {i}-{i + n - 1} -> {fig}")
        return
    prima = firma(doc, {p._p for i, n, *_ in piano for p in P[i:i + n]})
    shutil.copy2(DOC, BASE / "backup" / f"{DOC.stem}_{datetime.now():%Y%m%d_%H%M%S}{DOC.suffix}")
    for i, n, fig, larg, did in reversed(piano):          # dal fondo, cosi' gli indici restano validi
        ref = P[i]._p
        pf = doc.add_paragraph(); pf.alignment = WD_ALIGN_PARAGRAPH.CENTER; pf.paragraph_format.keep_with_next = True
        pf.add_run().add_picture(str(BASE / fig), width=Cm(larg))
        pc = doc.add_paragraph(); pc.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rc = pc.add_run(did); rc.italic = True; rc.font.size = Pt(9)
        ref.addprevious(pf._p); ref.addprevious(pc._p)
        nuovi += [pc._p]
        for p in P[i:i + n]:
            p._p.getparent().remove(p._p)
    dopo = firma(doc, set(nuovi))
    if dopo != prima:
        sys.exit("ERRORE: il resto del testo o le evidenziazioni sono cambiati: NON salvo.")
    doc.save(str(DOC))
    print(f"OK: {len(piano)} tavole sostituite con immagini; resto del testo ed evidenziazioni identici ({prima[1]} evidenziati)")


if __name__ == "__main__":
    main()
