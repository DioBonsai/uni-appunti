"""Integra il Word di Matematica discreta con contenuti del libro di Mori (16 ott 2026, su richiesta di Emanuele).
Regole: nessun paragrafo evidenziato viene modificato; si aggiungono paragrafi nuovi dopo un'ancora (inizio del testo del
paragrafo) e si modificano solo paragrafi non evidenziati. Dopo il salvataggio si controlla che testo ed evidenziazioni
dei paragrafi originali siano identici. Sintassi dei testi come in aggiungi_capitolo.py.
Uso: python integra_libro_mate.py [--dry-run]
"""
import shutil
import sys
from datetime import datetime
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parent))
from aggiungi_capitolo import (ORANGE, add_runs, doc_has_heading_styles, post_controlli, style_by_id)  # noqa: E402

UNI = Path(r"C:\Users\admin\Desktop\UNI")
WORD = UNI / "Riassunto Matematica discreta.docx"
BK = UNI / "agente" / "_sistema" / "backup"

# (indice di controllo, inizio del testo del paragrafo-ancora, [(tipo, testo), ...]); i nuovi paragrafi vanno DOPO l'ancora
INSERT = [
    (37, "Sottigliezza notazionale (⊆ vs ⊂)", [
        ("p", "**Attenzione ai testi.** Nel libro di testo e negli esercizi d'esame il simbolo ⊂ è usato con il significato di ⊆, cioè inclusione in senso lato, e la sua negazione ⊄ vuol dire semplicemente «non è sottoinsieme»; per esempio vi si trova «∅ ⊂ A» per ogni insieme A. In questi riassunti, invece, ⊂ indica l'inclusione stretta, come nella catena ℕ ⊂ ℤ ⊂ ℚ ⊂ ℝ. Davanti a un vero/falso conviene controllare sempre quale convenzione si sta usando: se ⊂ è letto come ⊆, l'affermazione «A ⊂ A» è vera."),
        ("h2", "Il quantificatore esistenziale e la negazione dei quantificatori"),
        ("p", "Accanto al quantificatore universale ∀ («per ogni») compare spesso il quantificatore esistenziale ∃ («esiste»): scrivere ∃ x ∈ A tale che P(x) significa che almeno un elemento di A soddisfa la proprietà P. Per dimostrare un'affermazione con ∃ basta esibire un elemento che la verifica; per smentire un'affermazione con ∀ basta esibire un controesempio."),
        ("p", "La negazione scambia i due quantificatori: ¬(∀ x ∈ A, P(x)) equivale a ∃ x ∈ A tale che ¬P(x), e ¬(∃ x ∈ A tale che P(x)) equivale a ∀ x ∈ A, ¬P(x). Per esempio, «A non è sottoinsieme di B» si scrive ∃ x ∈ A tale che x ∉ B: basta un elemento di A che non sta in B."),
    ]),
    (86, "Poiché la catena di implicazioni è valida in entrambe", [
        ("p", "**Distributività.** La doppia inclusione serve anche per la distributività. Per (X ∪ Y) ∩ Z = (X ∩ Z) ∪ (Y ∩ Z), si prende x ∈ (X ∪ Y) ∩ Z: allora x ∈ Z e x appartiene a X oppure a Y. Se x ∈ X, allora x ∈ X ∩ Z; se x ∈ Y, allora x ∈ Y ∩ Z; in entrambi i casi x ∈ (X ∩ Z) ∪ (Y ∩ Z), quindi il primo membro è contenuto nel secondo. Viceversa, se x ∈ (X ∩ Z) ∪ (Y ∩ Z), allora x ∈ X ∩ Z oppure x ∈ Y ∩ Z; in entrambi i casi x ∈ Z e x appartiene a X oppure a Y, cioè x ∈ (X ∪ Y) ∩ Z. Le due inclusioni danno l'uguaglianza."),
        ("h2", "Unione e intersezione di una famiglia di insiemi"),
        ("p", "Unione e intersezione si estendono a una famiglia qualunque di insiemi {A_{i}}_{i ∈ I}: l'unione è l'insieme degli x che appartengono ad almeno uno degli A_{i}, l'intersezione è l'insieme degli x che appartengono a tutti gli A_{i}."),
        ("p", "⋃_{i ∈ I} A_{i} = {x | ∃ i ∈ I, x ∈ A_{i}}   e   ⋂_{i ∈ I} A_{i} = {x | ∀ i ∈ I, x ∈ A_{i}}"),
        ("p", "Un esempio con gli intervalli aperti di ℝ: per n ≥ 1 sia A_{n} = (−1/n, 1/n). L'unione di tutti gli A_{n} è A₁ = (−1, 1), perché ogni A_{n} è contenuto in A₁; l'intersezione è {0}, perché un numero x ≠ 0 esce da A_{n} appena 1/n ≤ |x|. Con B_{n} = (−n, n) si ha invece ⋃ B_{n} = ℝ e ⋂ B_{n} = B₁ = (−1, 1)."),
        ("p", "**Notazione degli intervalli.** Per a < b si pone (a, b) = {x ∈ ℝ | a < x < b} (intervallo aperto), [a, b] = {x ∈ ℝ | a ≤ x ≤ b} (intervallo chiuso), e analogamente gli intervalli semiaperti [a, b) e (a, b]. Le semirette sono (a, +∞) = {x ∈ ℝ | x > a} e [0, +∞) = {x ∈ ℝ | x ≥ 0}, indicato anche con ℝ^{≥0}; l'insieme dei reali positivi si indica con ℝ^{>0}."),
    ]),
    (98, "Il prodotto cartesiano A × B come griglia di coppie ordinate.", [
        ("p", "Se uno dei due insiemi è vuoto, anche il prodotto è vuoto: ∅ × B = A × ∅ = ∅, perché non c'è modo di formare una coppia. Il prodotto si estende a n insiemi: A₁ × … × A_{n} è l'insieme delle n-uple ordinate (a₁, …, a_{n}) con a_{i} ∈ A_{i}, e se gli insiemi sono finiti |A₁ × … × A_{n}| = |A₁| · … · |A_{n}| (come si vedrà nel capitolo 22). La formula |A × B| = |A| · |B| vale anche quando un insieme è infinito, purché l'altro non sia vuoto: in quel caso il prodotto è infinito."),
    ]),
    (111, "Il complementare di S rispetto all'universo A.", [
        ("p", "**Notazione del libro.** Nel libro il complementare si scrive 𝒞_{X}(A), con X l'universo e A il sottoinsieme: le lettere sono scambiate rispetto a quanto scritto qui, ma la definizione è la stessa. Due esempi: 𝒞_{ℝ}(ℚ) = ℝ ∖ ℚ è l'insieme dei numeri irrazionali, e ℝ ∖ ℤ è il complementare di ℤ in ℝ. Differenza e complementare non vanno confusi: la differenza A ∖ B si può calcolare per due insiemi qualunque (anche quando B non è contenuto in A, e in tal caso A ∖ B = A ∖ (A ∩ B)), mentre il complementare presuppone che il sottoinsieme sia contenuto nell'universo."),
    ]),
    (144, " ⟺ x ∈ 𝒞A(S) ∪ 𝒞A(T)    (per definizione di unione)", [
        ("p", "Le leggi di De Morgan valgono per una famiglia qualunque di sottoinsiemi {S_{i}}_{i ∈ I} di A: 𝒞_{A}(⋃ S_{i}) = ⋂ 𝒞_{A}(S_{i}) e 𝒞_{A}(⋂ S_{i}) = ⋃ 𝒞_{A}(S_{i}). La dimostrazione è la stessa, con i quantificatori ∃ e ∀ al posto di ∨ e ∧ (la negazione di ∃ è ∀ e viceversa). Esiste anche la versione con la differenza, valida per insiemi qualunque: X ∖ (S ∪ T) = (X ∖ S) ∩ (X ∖ T) e X ∖ (S ∩ T) = (X ∖ S) ∪ (X ∖ T)."),
    ]),
    (186, "L'effetto domino: la base fa cadere la prima tessera", [
        ("p", "**Perché il principio funziona.** Si deduce dal quinto assioma. Sia P(n) una proprietà e sia U = {n ∈ ℕ | P(n) è vera}. Se vale la base P(0), allora 0 ∈ U; se vale il passo induttivo, per ogni n ∈ U si ha S(n) ∈ U, cioè anche n + 1 appartiene a U. Quindi U soddisfa le due condizioni dell'assioma di induzione e U = ℕ: la proprietà è vera per ogni n. Se la base è n = 1 si applica lo stesso ragionamento a U = {n ∈ ℕ | P(n + 1)}, e più in generale la base può essere un qualunque n₀: la proprietà risulta dimostrata per ogni n ≥ n₀."),
    ]),
    (199, "La formula per n+1 è verificata e la proprietà vale per ogni n ≥ 1.", [
        ("p", "**Un'induzione con base diversa da 0 e da 1.** Si dimostra che n² > 2n + 1 per ogni n ≥ 3. Base (n = 3): 9 > 7. Passo induttivo: si assume n² > 2n + 1 e si somma 2n + 1 a entrambi i membri, ottenendo (n + 1)² = n² + 2n + 1 > (2n + 1) + (2n + 1) = 4n + 2 ≥ 2n + 3 = 2(n + 1) + 1, dove l'ultima disuguaglianza vale perché 2n ≥ 1. Analogamente 2^{n} > n² per ogni n ≥ 5: base 32 > 25; passo: 2^{n+1} = 2 · 2^{n} > 2n² = n² + n² > n² + 2n + 1 = (n + 1)², usando la disuguaglianza appena dimostrata."),
    ]),
    (233, "I singoli sottoinsiemi Ai che compongono la partizione", [
        ("p", "**Nota di notazione.** In questo capitolo e nel successivo la partizione è indicata con 𝒫 = {A_{i}}, ma lo stesso simbolo 𝒫 indica anche l'insieme delle parti 𝒫(S). Per evitare confusione il libro indica con 𝒜 la famiglia di sottoinsiemi (ricoprimento o partizione) e riserva 𝒫(S) all'insieme delle parti: quando le due cose compaiono nella stessa formula, come in Q ⊆ 𝒫(S), la 𝒫 senza argomento va letta come 𝒜."),
    ]),
    (255, "Sebbene l'unione dia Y", [
        ("h2", "Altri esempi sulle partizioni"),
        ("p", "Dato un insieme S e un suo sottoinsieme A con A ≠ ∅ e A ≠ S, la famiglia {A, 𝒞_{S}(A)} è una partizione di S: i due blocchi non sono vuoti, sono disgiunti e la loro unione è S. In ℝ gli intervalli A_{n} = [n, n + 1] con n ∈ ℤ sono un ricoprimento ma non una partizione, perché due intervalli consecutivi si toccano in un estremo (il numero n + 1 appartiene sia ad A_{n} sia ad A_{n+1}); gli intervalli semiaperti [n, n + 1) formano invece una partizione di ℝ. Infine, per ogni intero n ≥ 1 sia P_{n} l'insieme delle frazioni m/n ridotte ai minimi termini con denominatore n: al variare di n i P_{n} sono una partizione di ℚ, perché ogni numero razionale si scrive in modo unico come frazione ridotta con denominatore positivo."),
    ]),
    (391, "Chi è abituato alle funzioni reali di variabile reale", [
        ("p", "Un esempio tipico: 1/x non definisce una funzione da ℝ in ℝ, perché in 0 non è definita; è una funzione da ℝ ∖ {0} in ℝ. Per averne una definita su tutto ℝ bisognerebbe assegnare comunque un valore in 0, e questo cambierebbe la funzione."),
    ]),
    (428, "Inclusione canonica j : S → A", [
        ("h2", "La funzione quoziente"),
        ("p", "Se ∼ è una relazione di equivalenza su X e Q = X/∼ è l'insieme quoziente, la **funzione quoziente** q : X → Q associa a ogni elemento la sua classe, q(x) = [x]. È suriettiva (ogni classe è la classe di un suo elemento) e q(x) = q(y) se e solo se x ∼ y: due elementi hanno la stessa immagine proprio quando sono equivalenti. Per una partizione, q manda ogni elemento nel blocco che lo contiene."),
    ]),
    (432, "Dunque una successione non è altro che una funzione", [
        ("p", "Una precisazione: se la successione parte da 0 il dominio è ℕ, se parte da 1 il dominio è ℕ_{≥1} = {1, 2, 3, …}. Formalmente sono funzioni diverse, perché cambia il dominio, anche se la lista dei valori è la stessa a meno di uno spostamento di indice; in entrambi i casi si parla di successione."),
    ]),
    (458, "La regola F([x]) = [2x] manda la stessa classe", [
        ("p", "**Un esempio in cui la regola funziona.** Sia ℤ diviso in pari e dispari, Q = {2ℤ, 2ℤ + 1} (con 2ℤ l'insieme dei pari e 2ℤ + 1 quello dei dispari), e si definisca G : Q → {−1, 1} con la regola G([x]) = (−1)^{x}. Se x ∼ y, cioè x e y hanno la stessa parità, allora (−1)^{x} = (−1)^{y}: il risultato dipende solo dalla classe e non dal rappresentante, quindi G è ben definita (vale 1 sui pari e −1 sui dispari). La differenza con F([x]) = [2x] sta tutta qui."),
    ]),
    (487, "La controimmagine di T = [0, 4] tramite x² è [−2, 2].", [
        ("p", "Altri esempi di controimmagine. Per la funzione costante f_{β} : A → B, la controimmagine f_{β}^{-1}(T) è tutto A se β ∈ T ed è ∅ altrimenti. Per la seconda proiezione p₂ : A × B → B si ha p₂^{-1}(T) = A × T. Per il valore assoluto |x| da ℝ in ℝ, la fibra di b > 0 è {−b, b}, la fibra di 0 è {0}, e la controimmagine di [0, 4] è [−4, 4]."),
    ]),
    (504, "Le tre situazioni: iniettiva, suriettiva, biettiva.", [
        ("p", "Il criterio con le controimmagini si dimostra subito: f non è iniettiva se e solo se esiste b ∈ B con |f^{-1}(b)| ≥ 2. Se f non è iniettiva ci sono a₁ ≠ a₂ con f(a₁) = f(a₂) = b, quindi a₁ e a₂ stanno entrambi in f^{-1}(b); viceversa due elementi distinti di f^{-1}(b) hanno la stessa immagine b. Un altro esempio di funzione biettiva è lo scambio A × B → B × A, (a, b) ↦ (b, a), la cui inversa è (b, a) ↦ (a, b). Il logaritmo log : ℝ^{>0} → ℝ è suriettivo (e iniettivo), mentre la funzione esponenziale da ℝ in ℝ non è suriettiva perché non assume valori minori o uguali a 0."),
    ]),
    (615, "Dunque f ∘ idA = f e idB ∘ f = f", [
        ("p", "**Composizione e funzioni costanti.** Sia f_{β} : A → B la funzione costante di valore β ∈ B e sia g : B → C una funzione qualunque. Allora g ∘ f_{β} è la funzione costante di valore g(β), cioè g ∘ f_{β} = f_{g(β)}, perché per ogni a si ha g(f_{β}(a)) = g(β). Se invece si compone con h : D → A, si ha f_{β} ∘ h = f_{β} (come funzione da D in B): ogni elemento viene comunque mandato in β."),
    ]),
    (683, "Lo stesso accade per l'esponenziale e il logaritmo", [
        ("h2", "Inversa di una composizione e cancellazione"),
        ("p", "**Proposizione.** Se f : A → B e g : B → C sono biettive, anche g ∘ f è biettiva e (g ∘ f)^{-1} = f^{-1} ∘ g^{-1}."),
        ("p", "**Dimostrazione.** g ∘ f è iniettiva e suriettiva per la Proposizione 1 del capitolo 19, quindi è biettiva. Per trovare l'inversa si controlla che f^{-1} ∘ g^{-1} : C → A funzioni: (f^{-1} ∘ g^{-1}) ∘ (g ∘ f) = f^{-1} ∘ (g^{-1} ∘ g) ∘ f = f^{-1} ∘ id_{B} ∘ f = f^{-1} ∘ f = id_{A}, e allo stesso modo (g ∘ f) ∘ (f^{-1} ∘ g^{-1}) = g ∘ (f ∘ f^{-1}) ∘ g^{-1} = g ∘ g^{-1} = id_{C}. Per l'unicità dell'inversa, f^{-1} ∘ g^{-1} è proprio (g ∘ f)^{-1}. L'ordine si inverte: per tornare indietro si annulla prima l'ultima operazione."),
        ("p", "**L'inversa dell'inversa.** Poiché f^{-1} ∘ f = id_{A} e f ∘ f^{-1} = id_{B}, la funzione f è a sua volta inversa di f^{-1}. Quindi f^{-1} è biettiva e (f^{-1})^{-1} = f."),
        ("p", "**Cancellazione.** Se f è suriettiva e g ∘ f = h ∘ f (con g, h : B → C), allora g = h: per ogni y ∈ B esiste x con f(x) = y, e g(y) = g(f(x)) = h(f(x)) = h(y). Se g è iniettiva e g ∘ f = g ∘ h (con f, h : A → B), allora f = h: per ogni a, da g(f(a)) = g(h(a)) segue f(a) = h(a). Le ipotesi non si possono togliere. Se f non è suriettiva può essere g ∘ f = h ∘ f con g ≠ h: per esempio f : {0} → {0, 1} con f(0) = 0, g l'identità di {0, 1} e h la funzione che manda tutto in 0. Se g non è iniettiva può essere g ∘ f = g ∘ h con f ≠ h: per esempio con g costante, perché g ∘ f e g ∘ h sono allora la stessa funzione costante qualunque siano f e h."),
    ]),
]

# (indice di controllo, testo da sostituire, nuovo testo): solo su paragrafi NON evidenziati
REPLACE = [
    (519, " (in generale basta che il codominio di f sia contenuto nel dominio di g)", ""),
]

# paragrafi vuoti (titolo vuoto accidentale) da togliere: (indice, tipo atteso)
DELETE_EMPTY = [(444, "Heading 1")]


def hl(p):
    return tuple(1 if r.font.highlight_color is not None else 0 for r in p.runs)


def firma(doc, escludi):
    return [(i, p.text, hl(p)) for i, p in enumerate(doc.paragraphs) if i not in escludi]


def replace_in_par(p, old, new):
    s = p.text.find(old)
    if s < 0:
        return False
    e = s + len(old)
    pos, first = 0, None
    for r in p.runs:
        rs, re_ = pos, pos + len(r.text)
        pos = re_
        if re_ <= s or rs >= e:
            continue
        a, b = max(s, rs) - rs, min(e, re_) - rs
        if first is None:
            first = r
            r.text = r.text[:a] + new + r.text[b:]
        else:
            r.text = r.text[b:]
        if e <= re_:
            break
    return first is not None


def main():
    dry = "--dry-run" in sys.argv
    lock = WORD.with_name("~$" + WORD.name[2:])
    if lock.exists() and not dry:
        sys.exit(f"Il Word e' aperto ({lock.name}). Chiudilo e riprova.")
    doc = Document(str(WORD))
    pars = list(doc.paragraphs)
    h2 = style_by_id(doc, "Titolo2")
    usa_h2 = h2 is not None and doc_has_heading_styles(doc)

    # 1) ancore e controlli, prima di toccare qualcosa
    ancore = []
    for idx, inizio, items in INSERT:
        p = pars[idx]
        if not p.text.startswith(inizio):
            sys.exit(f"ANCORA NON TROVATA al paragrafo {idx}: atteso «{inizio[:50]}», trovato «{p.text[:50]}»")
        ancore.append((p, items))
    for idx, old, new in REPLACE:
        if any(hl(pars[idx])) or old not in pars[idx].text:
            sys.exit(f"SOSTITUZIONE NON POSSIBILE al paragrafo {idx}")
    for idx, tipo in DELETE_EMPTY:
        p = pars[idx]
        if p.text.strip() or p.style.name != tipo or p._p.xpath(".//w:drawing"):
            sys.exit(f"Il paragrafo {idx} non e' un titolo vuoto")
    escludi = {idx for idx, _, _ in REPLACE} | {idx for idx, _ in DELETE_EMPTY}
    prima = firma(doc, escludi)

    # 2) modifiche
    n_new = 0
    creati = set()
    for p_anchor, items in ancore:
        cursore = p_anchor._p
        for tipo, testo in items:
            np_ = doc.add_paragraph()           # creato in fondo, poi spostato dopo l'ancora
            if tipo == "h2":
                add_runs(np_, testo, color=ORANGE)
                if usa_h2:
                    np_.style = h2
            else:
                add_runs(np_, testo)
            creati.add(np_.text)
            cursore.addnext(np_._p)
            cursore = np_._p
            n_new += 1
    for idx, old, new in REPLACE:
        replace_in_par(pars[idx], old, new)
    testi_mod = {pars[idx].text for idx, _, _ in REPLACE}   # testo dopo la sostituzione
    for idx, _ in DELETE_EMPTY:
        pars[idx]._p.getparent().remove(pars[idx]._p)

    if dry:
        print(f"[dry-run] aggiungerei {n_new} paragrafi, {len(REPLACE)} sostituzioni, {len(DELETE_EMPTY)} titolo/i vuoto/i tolti")
        return
    BK.mkdir(exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    shutil.copy2(WORD, BK / f"{WORD.stem}_{stamp}{WORD.suffix}")
    doc.save(str(WORD))

    # 3) controllo: i paragrafi originali (esclusi quelli modificati) hanno stesso testo ed evidenziazioni
    dopo_doc = Document(str(WORD))
    nuovi_testi = creati
    dopo = [(p.text, hl(p)) for p in dopo_doc.paragraphs if p.text not in nuovi_testi and p.text not in testi_mod]
    atteso = [(t, h) for _, t, h in prima]
    # i paragrafi sostituiti/cancellati sono stati tolti da 'atteso'; la sostituzione e' sul 519 (non evidenziato)
    dopo_senza, atteso_senza = dopo, atteso
    # un titolo vuoto in meno: tolgo i paragrafi vuoti Heading dal confronto non serve, e' gia' escluso
    if dopo_senza != atteso_senza:
        print(len(dopo_senza), len(atteso_senza))
        for k, (x, y) in enumerate(zip(dopo_senza, atteso_senza)):
            if x != y:
                print("prima differenza", k, repr(x[0][:80]), x[1], "|", repr(y[0][:80]), y[1]); break
        print("ATTENZIONE: testo o evidenziazioni dei paragrafi originali differiscono! Ripristino dal backup.")
        shutil.copy2(BK / f"{WORD.stem}_{stamp}{WORD.suffix}", WORD)
        sys.exit(1)
    print(f"OK: {n_new} paragrafi aggiunti, testo ed evidenziazioni originali identici (backup in _sistema/backup)")
    post_controlli(WORD)


if __name__ == "__main__":
    main()
