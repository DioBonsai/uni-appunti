"""Inserisce nei Word i paragrafi «In parole povere: ...» dopo le definizioni (regola di Emanuele, 19 ott 2026).
I testi stanno in _sistema/parole_povere/*.json: lista di {idx, ancora, testo}; il paragrafo nuovo va DOPO l'ancora.
Nessun paragrafo esistente viene modificato (evidenziati compresi): si controlla che testo ed evidenziazioni originali
restino identici. Uso: python applica_parole_povere.py [--dry-run] [materia ...]   (materie: algebra mate fondamenti programmazione laboratorio)
"""
import json
import shutil
import sys
from datetime import datetime
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parent))
from aggiungi_capitolo import add_runs, post_controlli  # noqa: E402

UNI = Path(__file__).resolve().parents[3]   # cartella UNI (accanto a agente/)
BASE = UNI / "agente" / "_sistema"
PP = BASE / "parole_povere"
BK = BASE / "backup"
GRUPPI = {
    "algebra": ("Riassunto Algebra lineare.docx", ["algebra.json"]),
    "mate": ("Riassunto Matematica discreta.docx", ["mate_1_12.json", "mate_13_25.json", "mate_extra.json"]),
    "fondamenti": ("Riassunto Fondamenti Informatica.docx", ["fondamenti.json"]),
    "programmazione": ("Riassunto Programmazione.docx", ["programmazione.json"]),
    "laboratorio": ("Riassunto Laboratorio.docx", ["laboratorio.json"]),
}


def hl(p):
    return tuple(1 if r.font.highlight_color is not None else 0 for r in p.runs)


def applica(nome, dry):
    docx_name, files = GRUPPI[nome]
    W = UNI / docx_name
    lock = W.with_name("~$" + W.name[2:])
    if lock.exists() and not dry:
        print(f"[{nome}] Word aperto ({lock.name}): salto")
        return
    voci = []
    for f in files:
        p = PP / f
        if p.exists():
            voci += json.loads(p.read_text(encoding="utf-8"))
        else:
            print(f"[{nome}] manca {f}")
    if not voci:
        return
    doc = Document(str(W))
    pars = list(doc.paragraphs)
    prima = [(p.text, hl(p)) for p in pars]
    per_ancora, scartate = {}, []
    for v in voci:
        idx, anc = v["idx"], v["ancora"]
        if 0 <= idx < len(pars) and pars[idx].text.startswith(anc[:50]):
            k = idx
        else:
            cand = [i for i, p in enumerate(pars) if p.text.startswith(anc[:50])]
            if len(cand) != 1:
                scartate.append((idx, anc[:40], len(cand)))
                continue
            k = cand[0]
        per_ancora.setdefault(k, []).append(v["testo"].strip())
    creati = set()
    n = 0
    for k, testi in per_ancora.items():
        cursore = pars[k]._p
        for t in testi:
            np_ = doc.add_paragraph()
            if t.startswith("In parole povere:"):
                t = "**In parole povere:**" + t[len("In parole povere:"):]
            add_runs(np_, t)
            creati.add(np_.text)
            cursore.addnext(np_._p)
            cursore = np_._p
            n += 1
    if scartate:
        print(f"[{nome}] ancore non trovate: {scartate}")
    if dry:
        print(f"[{nome}] [dry-run] aggiungerei {n} paragrafi")
        return
    BK.mkdir(exist_ok=True)
    shutil.copy2(W, BK / f"{W.stem}_prima-parole-povere_{datetime.now():%Y%m%d_%H%M%S}{W.suffix}")
    doc.save(str(W))
    dopo_doc = Document(str(W))
    dopo = [(p.text, hl(p)) for p in dopo_doc.paragraphs if p.text not in creati]
    if dopo != prima:
        print(f"[{nome}] ATTENZIONE: i paragrafi originali differiscono! Ripristino dal backup.")
        ultimo = sorted(BK.glob(f"{W.stem}_prima-parole-povere_*"))[-1]
        shutil.copy2(ultimo, W)
        return
    print(f"[{nome}] OK: {n} paragrafi aggiunti, testo ed evidenziazioni originali identici")


def main():
    dry = "--dry-run" in sys.argv
    scelte = [a for a in sys.argv[1:] if not a.startswith("--")] or list(GRUPPI)
    for nome in scelte:
        applica(nome, dry)
    if not dry:
        post_controlli(UNI / GRUPPI[scelte[0]][0])


if __name__ == "__main__":
    main()
