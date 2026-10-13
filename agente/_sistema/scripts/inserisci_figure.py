"""Inserisce nel Word le figure di un piano (piano.json) DOPO i paragrafi indicati, senza toccare testo ed evidenziazioni.

Uso:
    python inserisci_figure.py "Riassunto Matematica discreta.docx" figure/mate/piano.json [--dry-run] [--copia DEST.docx]

piano.json: lista di {dopo: indice paragrafo (python-docx), figura: percorso relativo a _sistema/, didascalia, larghezza (cm)}.
Sicurezza: backup in _sistema/backup/, rifiuta se il Word e' aperto, confronta testi ed evidenziazioni prima/dopo
(se qualcosa cambia non salva), poi verifica_word.py e backup_onedrive.py.
"""
import argparse
import json
import shutil
import sys
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

BASE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(Path(__file__).resolve().parent))
from percorsi_word import trova  # noqa: E402
from aggiungi_capitolo import post_controlli  # noqa: E402
UNI = BASE.parent.parent


def firma(doc, escludi=()):
    """(testi dei paragrafi, numero di run evidenziati) - deve restare identico, tolte le figure."""
    testi, evid = [], 0
    for p in doc.paragraphs:
        if p._p.xpath(".//w:drawing") or p._p in escludi:
            continue
        testi.append(p.text)
        for r in p.runs:
            rpr = r._r.rPr
            if rpr is not None and (rpr.find(qn("w:highlight")) is not None or rpr.find(qn("w:shd")) is not None):
                evid += 1
    return testi, evid


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("docx"); ap.add_argument("piano")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--copia", help="salva su questo file invece di modificare l'originale (prova)")
    a = ap.parse_args()

    target = Path(a.docx)
    if not target.is_absolute():
        target = UNI / target
    if not target.exists():
        target = trova(UNI, target.parent.name, target.name)
    piano_p = Path(a.piano)
    if not piano_p.is_absolute():
        piano_p = BASE / piano_p
    piano = json.loads(piano_p.read_text(encoding="utf-8"))

    lock = target.with_name("~$" + target.name[2:])
    if lock.exists() and not a.copia and not a.dry_run:
        sys.exit(f"Il Word e' aperto ({lock.name}). Chiudilo e riprova.")

    doc = Document(str(target))
    pars = doc.paragraphs
    if any(p._p.xpath(".//w:drawing") for p in pars):
        print("ATTENZIONE: il Word contiene gia' immagini; controllo che gli indici del piano siano ancora validi.")
    prima = firma(doc)
    for f in piano:
        if not (0 <= f["dopo"] < len(pars)):
            sys.exit(f"Indice {f['dopo']} fuori dal documento ({len(pars)} paragrafi)")
        if not (BASE / f["figura"]).exists():
            sys.exit(f"Figura mancante: {f['figura']}")

    if a.dry_run:
        for f in piano:
            print(f"dopo [{f['dopo']}] {pars[f['dopo']].text[:70]!r}  <-  {Path(f['figura']).name}")
        print(f"[dry-run] inserirei {len(piano)} figure in {target.name}")
        return

    didasc = []
    # dal fondo verso l'alto, cosi' gli indici dei paragrafi precedenti restano validi
    for f in sorted(piano, key=lambda x: x["dopo"], reverse=True):
        ref = pars[f["dopo"]]._p
        pc = None
        if f.get("didascalia"):
            pc = doc.add_paragraph(); pc.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rc = pc.add_run(f["didascalia"]); rc.italic = True; rc.font.size = Pt(9); didasc.append(pc._p)
        pf = doc.add_paragraph(); pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.paragraph_format.keep_with_next = True
        pf.paragraph_format.space_before = Pt(6)
        pf.add_run().add_picture(str(BASE / f["figura"]), width=Cm(float(f.get("larghezza", 11))))
        ref.addnext(pf._p)
        if pc is not None:
            pf._p.addnext(pc._p)

    dopo = firma(doc, didasc)
    if dopo != prima:
        sys.exit("ERRORE: testo o evidenziazioni sono cambiati: NON salvo.")
    n = len(doc.paragraphs)

    if a.copia:
        out = Path(a.copia); doc.save(str(out))
        print(f"OK (copia di prova): {len(piano)} figure -> {out} ({n} paragrafi); testo ed evidenziazioni identici")
        return
    bk = BASE / "backup"; bk.mkdir(exist_ok=True)
    shutil.copy2(target, bk / f"{target.stem}_{datetime.now():%Y%m%d_%H%M%S}{target.suffix}")
    doc.save(str(target))
    print(f"OK: {len(piano)} figure inserite in {target.name}; testo ed evidenziazioni identici ({prima[1]} evidenziati)")
    post_controlli(target)


if __name__ == "__main__":
    main()
