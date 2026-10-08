"""Applica le correzioni degli audit (8 ott 2026) ai quattro Word, con copia di sicurezza prima.
Ogni modifica e' una sostituzione di testo mirata (le evidenziazioni restano dove sono); le sostituzioni che non trovano
il testo vengono segnalate come MANCA senza toccare nulla. Eseguito una sola volta: tenuto come registro delle modifiche.
"""
import shutil
import sys
from datetime import datetime
from pathlib import Path

from docx import Document

UNI = Path(r"C:\Users\admin\Desktop\UNI")
BK = UNI / "agente" / "_sistema" / "backup"
ts = datetime.now().strftime("%Y%m%d_%H%M%S")
report = []


def variants(s):
    return [s, s.replace("'", "\u2019"), s.replace('"', "\u201c").replace("'", "\u2019")]


def find_par(d, old):
    for v in variants(old):
        for p in d.paragraphs:
            if v in p.text:
                return p, v
    return None, None


def replace_in_par(p, old, new):
    text = p.text
    s = text.find(old)
    if s < 0:
        return False
    e = s + len(old)
    pos, first = 0, None
    for r in p.runs:
        rs, re_ = pos, pos + len(r.text)
        pos = re_
        if re_ <= s or rs >= e:
            continue
        a = max(s, rs) - rs
        b = min(e, re_) - rs
        if first is None:
            first = r
            r.text = r.text[:a] + new + r.text[b:]
        else:
            r.text = r.text[b:]
        if e <= re_:
            break
    return first is not None


def edit(d, label, old, new):
    p, v = find_par(d, old)
    if p is None:
        report.append(("MANCA", label))
        return
    ok = replace_in_par(p, v, new)
    report.append(("OK   " if ok else "MANCA", label))


def append(d, label, anchor, text):
    p, v = find_par(d, anchor)
    if p is None:
        report.append(("MANCA", label))
        return
    p.add_run(text)
    report.append(("OK   ", label))


def delete_par(d, label, anchor):
    p, v = find_par(d, anchor)
    if p is None:
        report.append(("MANCA", label))
        return
    p._p.getparent().remove(p._p)
    report.append(("OK   ", label))


def run_doc(materia, fname, ops):
    path = UNI / materia / fname
    shutil.copy2(path, BK / f"{path.stem}_prima-audit_{ts}{path.suffix}")
    d = Document(str(path))
    report.append(("----", materia))
    ops(d)
    d.save(str(path))


# ====================================================================== MATEMATICA DISCRETA
def mate(d):
    edit(d, "cap1: formula del sottoinsieme", "S ⊆ T ⟺ (∀ x ∈ S ⟹ x ∈ T)", "S ⊆ T ⟺ ∀ x ∈ S, x ∈ T")
    edit(d, "cap1: tolto «(o impropri)»", "(o impropri)", "")
    edit(d, "cap1: sottoinsieme proprio coerente con la definizione", "qualsiasi altro sottoinsieme intermedio è detto proprio",
         "un sottoinsieme S ⊆ A con S ≠ A è detto proprio (come definito sopra)")
    edit(d, "cap2: elementi non «ereditati»",
         "gli elementi contenuti all'interno di un sotto-insieme non vengono ereditati automaticamente dall'insieme contenitore di livello superiore.",
         "gli elementi di un insieme che è a sua volta elemento di B (come A in B = {A, 3}) non sono automaticamente elementi di B: A ∈ B non implica A ⊆ B.")
    edit(d, "cap4: differenza non commutativa",
         "la differenza tra insiemi non gode della proprietà commutativa, ossia A ∖ B ≠ B ∖ A.",
         "la differenza tra insiemi non è commutativa: esistono insiemi per cui A ∖ B ≠ B ∖ A (ad esempio per A = {1, 2} e B = {2, 3} si ha A ∖ B = {1} e B ∖ A = {3}).")
    edit(d, "cap4: tolto «oppure Sc»", " oppure Sc", "")
    edit(d, "cap4: «triletto» → insieme intero", "il triletto {a, b, c}", "l'insieme intero {a, b, c}")
    edit(d, "cap5: Peano e Russell (causalità storica)",
         "Tuttavia, la scoperta di paradossi logici fondamentali (come il Paradosso di Russell sulla collezione di tutti gli insiemi che non contengono se stessi) fece crollare i tentativi di costruzione diretta e dimostrò la necessità di definire l'aritmetica mediante un sistema formale di assiomi.",
         "Gli assiomi di Peano (1889) seguono un'altra strada: non costruiscono i numeri naturali a partire dagli insiemi, ma li caratterizzano dando un elemento iniziale, lo 0, e una funzione successore con proprietà precise. (Il paradosso di Russell, del 1901, relativo alla collezione di tutti gli insiemi che non contengono se stessi, mostrò poi che la teoria ingenua degli insiemi va maneggiata con cautela.)")
    edit(d, "cap5: esempio sbagliato sul quinto assioma",
         "ad esempio, una sequenza infinita per i numeri pari seguita da una sequenza isolata per i dispari, da cui non si potrebbe mai raggiungere la seconda partendo da 0",
         "ad esempio, ℕ affiancato da una seconda catena infinita separata (…, −1′, 0′, 1′, …) i cui elementi non si raggiungono mai partendo da 0 con il successore: una struttura del genere soddisfa i primi quattro assiomi ma non il quinto")
    edit(d, "cap5: doppia parentesi", "(Raccoglimento parziale di ((n+1)))", "(Raccoglimento parziale di (n+1))")
    edit(d, "cap5: induzione parte da n = 1",
         "La formula per n+1 è verificata e la proprietà vale per ogni n ∈ ℕ.",
         "La formula per n+1 è verificata e la proprietà vale per ogni n ≥ 1.")


# ====================================================================== ALGEBRA
def alg(d):
    edit(d, "cap1: verso (non direzione)", "invertire la direzione del vettore", "invertire il verso del vettore")
    edit(d, "cap1: «struttura portante dell'informatica»", "l'algebra lineare costituisce la struttura portante dell'informatica",
         "l'algebra lineare è fondamentale in informatica")
    edit(d, "cap1: elaborazione «prevalentemente» mappe lineari",
         "e il processo di elaborazione dei dati consiste prevalentemente nell'applicazione di mappe lineari.",
         "e molte elaborazioni dei dati sono composte da mappe lineari.")
    edit(d, "cap1: reti neurali non sono solo mappe lineari",
         "appresi dal sistema durante l'addestramento non sono altro che trasformazioni e mappe lineari.",
         "appresi durante l'addestramento sono i coefficienti di mappe lineari, alternate a funzioni non lineari: l'algebra lineare è quindi una componente fondamentale di questi modelli.")
    edit(d, "cap1: pulizia del segnale «spesso» (1)", "L'operazione necessaria per", "Per")
    edit(d, "cap1: pulizia del segnale «spesso» (2)", "si realizza proprio applicando un'opportuna mappa lineare.",
         "si usa spesso un'opportuna mappa lineare.")
    edit(d, "cap2: da dove nascono gli insiemi numerici",
         "dove ciascun passaggio nasce dall'esigenza algebrica di risolvere un problema o un'operatore non consentito nell'insieme precedente.",
         "dove ℤ nasce per poter sottrarre, ℚ per poter dividere e ℝ per colmare i \"buchi\" dei razionali (completezza); ℂ, come vedremo, per risolvere equazioni prive di soluzioni reali.")
    edit(d, "cap2: Pitagora «secondo la tradizione»",
         "già nell'antica Grecia con la scuola pitagorica si scoprì che la diagonale di un quadrato di lato 1 ha lunghezza pari a √2.",
         "secondo la tradizione, già nell'antica Grecia un membro della scuola pitagorica si accorse che la diagonale di un quadrato di lato 1, lunga √2, non è commensurabile con il lato.")
    edit(d, "cap2: etichetta della distributiva", "Proprietà di Anello/Legame:", "Proprietà distributiva (lega somma e prodotto):")
    edit(d, "cap2: riferimento «proprietà 4» senza numerazione", " (proprietà 4)", "")
    edit(d, "cap2: campo degli scalari", "sarà quasi sempre ℝ (o ℂ per la parte sui numeri complessi).",
         "sarà generalmente ℝ oppure ℂ (la definizione vale però per qualsiasi campo).")
    edit(d, "cap3: ampliamenti numerici",
         "ogni ampliamento nasce dall'esigenza algebrica di risolvere problemi ed equazioni che non trovano soluzione nell'insieme precedente.",
         "ogni ampliamento nasce da un'esigenza precisa: sottrarre (ℤ), dividere (ℚ), colmare i \"buchi\" dei razionali (ℝ), risolvere equazioni prive di soluzioni reali (ℂ).")
    edit(d, "cap3: parte immaginaria", "chiamati rispettivamente parte reale e coefficiente della parte immaginaria.",
         "chiamati rispettivamente parte reale (a = Re(z)) e parte immaginaria (b = Im(z)): la parte immaginaria è il numero reale b, non bi.")
    edit(d, "cap4: confronti tra complessi (inizio frase)",
         "Gli unici confronti ammessi tra numeri complessi riguardano la loro eventuale uguaglianza",
         "I numeri complessi come tali si possono solo confrontare per uguaglianza")
    append(d, "cap4: confronti tra complessi (aggiunta)", "I numeri complessi come tali si possono solo confrontare per uguaglianza",
           " Si possono invece ordinare grandezze reali associate, come i moduli (|z1| < |z2|) o le parti reali (Re(z1) < Re(z2)).")
    edit(d, "cap6: asse immaginario",
         "ospita i numeri immaginari puri, corrispondenti ai complessi con parte reale nulla (0 + bi).",
         "ospita i numeri della forma bi con b ∈ ℝ (immaginari puri per b ≠ 0, più l'origine).")


# ====================================================================== FONDAMENTI
def fond(d):
    edit(d, "cap7: somma binaria 11010 + 11011",
         "Risultato finale: 1010101 (corrispondente al valore decimale 53).",
         "Risultato finale: 110101 (corrispondente al valore decimale 53 = 26 + 27).")
    edit(d, "cap3: capacità del Blu-ray", "capacità superiori a 50 GB.",
         "capacità oltre cinque volte quella di un DVD (25 GB a strato singolo, fino a 50 GB a doppio strato).")
    edit(d, "cap4: origine dell'ASCII",
         "Nato negli anni '40 e '50 per standardizzare la comunicazione tra macchine diverse, il codice ",
         "negli anni '40 e '50 esistevano molti codici tra loro incompatibili; per standardizzare la comunicazione tra macchine diverse l'ANSI (American National Standards Institute) adottò il codice ")
    edit(d, "cap4: tolta frase ripetuta sull'ANSI",
         " Lo standard è stato sviluppato dall'ANSI (American National Standards Institute).", "")


# ====================================================================== PROGRAMMAZIONE
def prog(d):
    delete_par(d, "cap1: «Anni '70, transistor» (falso e non nelle slide)", "Anni '70: con l'invenzione del transistor")
    edit(d, "cap3: iterazione (frase troncata)", " o viene soddisfatta", ".")
    edit(d, "cap3: linguaggi di basso livello",
         "come il linguaggio macchina o l'assembly, espressi direttamente in sequenze di bit e strettamente legati",
         "come il linguaggio macchina (sequenze di bit) o l'assembly (mnemonici testuali che lo rappresentano), strettamente legati")
    edit(d, "cap4: refuso «scorte»", "venivano scorte e pilotate", "venivano scelte e pilotate")
    edit(d, "cap4: macchine cablate (inizio)",
         "I primi calcolatori automatici progettati fino alla prima metà del Novecento erano macchine cablate",
         "Queste prime macchine erano macchine cablate")
    edit(d, "cap4: macchine cablate (esempio)", "(ad esempio solo addizioni o tabelle balistiche)", "(tipicamente solo addizione e sottrazione)")
    edit(d, "cap4: Babbage",
         "Macchina Analitica di Charles Babbage (1840), che introdusse l'uso delle schede perforate per guidare il flusso di calcolo",
         "Macchina Analitica di Charles Babbage (descritta intorno al 1840), primo esempio di macchina di calcolo programmabile, con dati e istruzioni su schede perforate")
    edit(d, "cap4: EDVAC, Mauchly/Eckert e von Neumann",
         "Nel 1944-1945, nell'ambito del progetto EDVAC, il matematico John von Neumann formalizzò il celebre Modello di Von Neumann (o architettura EDVAC), introducendo tre intuizioni rivoluzionarie",
         "L'EDVAC, progettato nel 1944 da John Mauchly e J. Presper Eckert (già sviluppatori dell'ENIAC), introdusse tre intuizioni rivoluzionarie. Il matematico John von Neumann, consulente del progetto, descrisse e pubblicò per primo questa architettura nel 1945: da qui il celebre Modello di Von Neumann (o architettura EDVAC), che comprende")
    edit(d, "cap4: bit realizzato con relè, valvole, transistor",
         "mediante transistor a due stati stabili (latch), valvole o relè",
         "mediante dispositivi come relè, valvole o transistor, che realizzano circuiti a due stati stabili")
    edit(d, "cap5: assembly 1 a 1 (istruzioni, non righe)",
         "Tra il codice assembly e il codice macchina esiste una corrispondenza biunivoca (1 a 1): ogni riga di testo assembly equivale esattamente a una singola istruzione numerica della CPU.",
         "Ogni istruzione assembly (LOAD, ADD, STORE, …) corrisponde a una singola istruzione macchina; le direttive come ADDR A = 400 definiscono soltanto indirizzi simbolici e non producono codice macchina.")
    edit(d, "cap5: verbosità dell'assembly",
         "un algoritmo descrivibile in due righe concettuali richiede circa 10 istruzioni macchina atomiche.",
         "un algoritmo semplice come la moltiplicazione richiede 10 righe di assembly (e l'addizione 4).")
    edit(d, "cap5: nascita di C e Unix",
         "Progettato nel 1972 da Dennis Ritchie presso i Bell Labs per sviluppare il sistema operativo UNIX, il linguaggio C è diventato",
         "Unix nasce nel 1969 ai Bell Labs (Ken Thompson, PDP-7, scritto in assembly); per superare i limiti dell'assembly Dennis Ritchie progetta il linguaggio C (prima versione nel 1972) e Unix viene poi riscritto in C. Il linguaggio C è diventato")
    edit(d, "cap5: C strutturato e goto", "eliminando l'uso dei salti incondizionati grezzi.", "per cui non servono salti a numeri di riga.")
    edit(d, "cap6: dimensione dei tipi non «fissa»",
         "il quantitativo fisso di byte occupati nella RAM per memorizzare un dato di quel tipo (ad esempio 4 byte per un int, 8 byte per un double).",
         "quanti byte occupa nella RAM un dato di quel tipo (tipicamente 4 byte per un int e 8 per un double; dipende dalla piattaforma).")
    edit(d, "cap6: commento «//» davanti a una frase", "// tipo nativo (o primitivo/built-in) è", "Un tipo nativo (o primitivo/built-in) è")
    edit(d, "cap6: long dipende anche dal sistema operativo", "long (4 o 8 byte a seconda dell'architettura)",
         "long (4 o 8 byte a seconda di architettura e sistema operativo)")
    edit(d, "cap6: overflow di int «di solito»", "che fa \"ripartire\"", "che di solito fa \"ripartire\"")
    append(d, "cap6: overflow di int con segno (aggiunta)", "dal limite minimo negativo.",
           " Per gli interi con segno lo standard C lo lascia formalmente indefinito.")
    edit(d, "cap6: commento overflow nel codice", "OVERFLOW: x diventa -2147483648",
         "OVERFLOW (indefinito per gli int con segno): tipicamente x diventa -2147483648")
    edit(d, "cap6: sizeof a compile time",
         "L'operatore unario sizeof() permette di ricavare a runtime il quantitativo esatto di memoria occupato",
         "L'operatore unario sizeof permette di ricavare, a tempo di compilazione, il numero esatto di byte occupati")
    edit(d, "cap6: commento «Inizializzazione/Assegnamento»", "// Inizializzazione/Assegnamento successivo: il valore diventa 10",
         "// Assegnamento successivo: il valore diventa 10")
    edit(d, "cap7: Stack Pointer",
         "memorizza l'indirizzo dinamico dell'ultima posizione occupata sulla pila, identificando il punto esatto in cui allocare la variabile successiva.",
         "indica l'indirizzo dinamico a cui allocare la variabile successiva (la cima corrente dello stack) e si sposta man mano che si allocano variabili.")
    edit(d, "cap8: assembly, non codice macchina",
         "nelle seguenti istruzioni macchina eseguite dalla CPU:",
         "nelle seguenti istruzioni assembly (poi tradotte in codice macchina), eseguite dalla CPU:")
    edit(d, "cap8: refuso «un ampiezza»", "occupando un ampiezza", "occupando un'ampiezza")


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    run_doc("Matematica discreta", "Appunti matematica discreta.docx", mate)
    run_doc("Algebra lineare", "Riassunto Algebra lineare.docx", alg)
    run_doc("Fondamenti Info", "Fondamenti Informatica.docx", fond)
    run_doc("Programmazione", "Riassunto Programmazione.docx", prog)
    for lv, m in report:
        print(lv, m)
    n_ok = sum(1 for l, _ in report if l.startswith("OK"))
    n_no = sum(1 for l, _ in report if l == "MANCA")
    print(f"\nApplicate: {n_ok} | non trovate: {n_no}")


if __name__ == "__main__":
    main()
