"""Aggiunge un capitolo (scritto in un file .md semplice) in fondo a un Word di riassunti,
replicando la formattazione dei tuoi documenti.

Uso:
    python aggiungi_capitolo.py "Algebra lineare/Riassunto Algebra lineare.docx" capitolo.md
    python aggiungi_capitolo.py ... capitolo.md --dry-run     # non scrive, mostra solo l'esito

Formato del .md:
    # 7. Titolo del capitolo            -> grassetto rosso (come i tuoi capitoli)
    ## Sottotitolo                      -> grassetto arancione
    testo con **grassetto** e *corsivo*
    - elenco puntato
    1. elenco numerato (resta testo con numero)
    riga vuota = nuovo paragrafo

Sicurezza: prima di scrivere fa una copia in _sistema/backup/ e si rifiuta di lavorare se il
file e' aperto in Word (esiste il file di blocco ~$...).
"""
import argparse
import re
import shutil
import sys
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor

BASE = Path(__file__).resolve().parent.parent
UNI = BASE.parent
RED = RGBColor(0xEE, 0x00, 0x00)
ORANGE = RGBColor(0xE9, 0x71, 0x32)
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")


def bullet_numid(doc: Document):
    """Trova un numId di elenco puntato gia' presente nel documento."""
    try:
        numbering = doc.part.numbering_part.element
    except Exception:
        return None
    bullet_abs = set()
    for a in numbering.findall(qn("w:abstractNum")):
        lvl = a.find(qn("w:lvl"))
        fmt = lvl.find(qn("w:numFmt")) if lvl is not None else None
        if fmt is not None and fmt.get(qn("w:val")) == "bullet":
            bullet_abs.add(a.get(qn("w:abstractNumId")))
    for n in numbering.findall(qn("w:num")):
        ref = n.find(qn("w:abstractNumId"))
        if ref is not None and ref.get(qn("w:val")) in bullet_abs:
            return n.get(qn("w:numId"))
    return None


def add_runs(par, text, color=None, bold=False):
    for chunk in INLINE.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            run = par.add_run(chunk[2:-2]); run.bold = True
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            run = par.add_run(chunk[1:-1]); run.italic = True
        else:
            run = par.add_run(chunk)
            if bold:
                run.bold = True
        if color is not None:
            run.bold = True
            run.font.color.rgb = color


def set_bullet(par, numid):
    pPr = par._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), "0")
    nid = OxmlElement("w:numId"); nid.set(qn("w:val"), str(numid))
    numPr.append(ilvl); numPr.append(nid)
    pPr.append(numPr)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("docx"); ap.add_argument("md")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()

    target = Path(a.docx)
    if not target.is_absolute():
        target = UNI / target
    src = Path(a.md)
    if not target.exists() or not src.exists():
        sys.exit(f"File non trovato: {target if not target.exists() else src}")

    # Word crea "~$" + nome senza i primi 2 caratteri mentre il file e' aperto
    lock = target.with_name("~$" + target.name[2:])
    if lock.exists() and not a.dry_run:
        sys.exit(f"Il Word e' aperto ({lock.name}). Chiudilo e riprova: non scrivo per non rovinarlo.")

    doc = Document(str(target))
    numid = bullet_numid(doc)
    n_par = 0
    for raw in src.read_text(encoding="utf-8").splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        p = doc.add_paragraph()
        n_par += 1
        if line.startswith("## "):
            add_runs(p, line[3:].strip(), color=ORANGE)
        elif line.startswith("# "):
            add_runs(p, line[2:].strip(), color=RED)
        elif re.match(r"^\s*[-*] ", line):
            add_runs(p, re.sub(r"^\s*[-*] ", "", line))
            if numid:
                set_bullet(p, numid)
        else:
            add_runs(p, line.strip())

    if a.dry_run:
        print(f"[dry-run] aggiungerei {n_par} paragrafi a {target.name} (elenco puntato: {'si' if numid else 'no'})")
        return

    bk = BASE / "backup"
    bk.mkdir(exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    shutil.copy2(target, bk / f"{target.stem}_{stamp}{target.suffix}")
    doc.save(str(target))
    print(f"OK: {n_par} paragrafi aggiunti a {target.name} (backup in _sistema/backup)")


if __name__ == "__main__":
    main()
